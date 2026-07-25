"""Shared style and layout toolkit so all sheets read as one document.

v5 repeated formatting code in every sheet function and drifted. Everything
visual lives here now.

Design rules:
  * Colour carries meaning. Hue = modality, bar length = price, icon = difficulty,
    badge = hazard/privacy. v5 gave 32 categories 32 arbitrary pastels, which
    looked busy and told you nothing.
  * Scan-view and read-view are different sheets. Catalog rows are one or two
    lines; prose lives on the Cards sheet.
  * Every sheet carries the same hyperlinked nav bar in row 1.

openpyxl constraints respected here (verified against 3.1.5 source):
  * A native Table brings its OWN autofilter. Never also set ws.auto_filter on a
    table sheet — two filter definitions makes Excel repair the file.
  * Never merge cells inside a table/filter range, including its header row.
    All merges are confined to the banner rows above the table.
  * column_dimensions.group() DELETES the widths of every column after the
    first, so group before setting widths.
  * No sparklines and no pivot tables exist in openpyxl. Use DataBarRule and
    precomputed cross-tabs driven by SUMIFS.
"""
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------- palette
NAVY = "1F2A44"
NAVY_SOFT = "3A4A6B"
INK = "24242B"
MUTED = "7A7A85"
PAPER = "FBF9F4"
LINE = "D8D2C6"
ACCENT = "C8853A"
ACCENT_SOFT = "F3E2C7"
GOOD = "3F7D5A"
WARN = "B8862B"
BAD = "A8412F"

FONT = "Aptos Narrow"       # Excel 2024+ default; falls back gracefully
MONO = "Consolas"

_thin = Side(style="thin", color=LINE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
BORDER_B = Border(bottom=_thin)

# Hazard and privacy badges — short, high-contrast, meaningful at a glance.
HAZARD_BADGE = {
    "Mains": "⚡MAINS", "HighVoltage": "⚡HV", "HotSurface": "🔥HOT", "Laser": "☀LASER",
    "UV": "☀UV", "Ignition": "💥IGNITION", "Asphyxiant": "☁O2", "Toxic": "☠TOXIC",
    "HighCurrent": "⚡AMPS", "Radiation": "☢RAD", "Mechanical": "⚙PINCH",
}
PRIVACY_BADGE = {
    "None": "—", "Aggregate": "◔ aggregate",
    "Identifiable": "◑ identifiable", "Raw-imagery": "● raw A/V",
}
PRIVACY_FILL = {
    "None": "E7F0E9", "Aggregate": "EAF0F6",
    "Identifiable": "FBEEDD", "Raw-imagery": "F8DED7",
}


# ---------------------------------------------------------------- primitives

def sfont(size=10, bold=False, italic=False, color=INK, name=FONT):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


def cell(ws, row, col, value=None, *, size=10, bold=False, italic=False,
         color=INK, fill=None, wrap=True, halign="left", valign="top",
         border=True, name=FONT):
    c = ws.cell(row=row, column=col, value=value)
    c.font = sfont(size, bold, italic, color, name)
    c.alignment = Alignment(wrap_text=wrap, horizontal=halign, vertical=valign)
    if border:
        c.border = BORDER
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def rich(*parts):
    """CellRichText from (text, **fontkwargs) pairs — for readable prose cells."""
    blocks = []
    for p in parts:
        if isinstance(p, str):
            blocks.append(p)
        else:
            text, kw = p
            blocks.append(TextBlock(InlineFont(rFont=kw.pop("name", FONT), **kw), text))
    return CellRichText(blocks)


def set_widths(ws, widths, start=1):
    """Call AFTER any column grouping — group() destroys widths."""
    for i, w in enumerate(widths, start):
        ws.column_dimensions[get_column_letter(i)].width = w


def link_to(c, sheet, anchor="A1", text=None, tooltip=None):
    if text is not None:
        c.value = text
    c.hyperlink = Hyperlink(ref=c.coordinate,
                            location=f"{quote_sheetname(sheet)}!{anchor}",
                            display=str(c.value), tooltip=tooltip)
    c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold,
                  color="1F4E79", underline="single")
    return c


# ---------------------------------------------------------------- page furniture

def nav_bar(ws, sheets, current, ncols):
    """Row 1: hyperlinked navigation, identical on every sheet."""
    ws.row_dimensions[1].height = 17
    col = 1
    for name, label in sheets:
        if col > ncols:
            break
        c = ws.cell(row=1, column=col, value=label)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if name == current:
            c.font = sfont(8, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=ACCENT)
        else:
            c.font = Font(name=FONT, size=8, color="FFFFFF", underline=None)
            c.fill = PatternFill("solid", fgColor=NAVY_SOFT)
            c.hyperlink = Hyperlink(ref=c.coordinate,
                                    location=f"{quote_sheetname(name)}!A1",
                                    display=label, tooltip=name)
        col += 1
    for c2 in range(col, ncols + 1):
        k = ws.cell(row=1, column=c2)
        k.fill = PatternFill("solid", fgColor=NAVY_SOFT)


def banner(ws, title, subtitle, ncols, row=2):
    """Title block. Merges live here, strictly ABOVE any table/filter range."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    t = ws.cell(row=row, column=1, value=title)
    t.font = Font(name=FONT, size=17, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 32
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=NAVY)

    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
    s = ws.cell(row=row + 1, column=1, value=subtitle)
    s.font = sfont(9.5, italic=True, color=MUTED)
    s.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=True)
    ws.row_dimensions[row + 1].height = 26
    return row + 2


def section(ws, row, text, ncols, note=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = sfont(12, bold=True, color=NAVY)
    c.fill = PatternFill("solid", fgColor=ACCENT_SOFT)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 24
    row += 1
    if note:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        n = ws.cell(row=row, column=1, value=note)
        n.font = sfont(9, italic=True, color=MUTED)
        n.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        ws.row_dimensions[row].height = 22
        row += 1
    return row


def table_header(ws, row, headers, fill=NAVY):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = sfont(9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER


def make_table(ws, first_row, last_row, ncols, name, style="TableStyleLight15"):
    """Native Excel Table. Brings its own filter — never set ws.auto_filter too.
    Header cells must already be written and must be unique strings."""
    if last_row <= first_row:
        return None
    ref = f"A{first_row}:{get_column_letter(ncols)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True,
                                      showColumnStripes=False,
                                      showFirstColumn=False, showLastColumn=False)
    ws.add_table(t)
    return t


def unique_headers(headers):
    """Excel repairs the file on duplicate table column names; openpyxl won't stop you."""
    seen, out = {}, []
    for h in headers:
        h = str(h) if h not in (None, "") else "col"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 0
        out.append(h)
    return out


# ---------------------------------------------------------------- conditional formatting

def price_bars(ws, col, first, last, max_price=120):
    rng = f"{get_column_letter(col)}{first}:{get_column_letter(col)}{last}"
    ws.conditional_formatting.add(rng, DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=max_price,
        color="9BB7D4", showValue=True))


def difficulty_icons(ws, col, first, last):
    """Numeric 1-5 difficulty rendered as a 5-step icon ramp."""
    rng = f"{get_column_letter(col)}{first}:{get_column_letter(col)}{last}"
    ws.conditional_formatting.add(rng, IconSetRule(
        "5Rating", "num", [1, 2, 3, 4, 5], showValue=True, reverse=False))


def flag_text(ws, first_col, last_col, first, last, formula, color=BAD, fill="FBE3DE"):
    rng = f"{get_column_letter(first_col)}{first}:{get_column_letter(last_col)}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[formula], font=sfont(9, bold=True, color=color),
        fill=PatternFill("solid", fgColor=fill)))


def highlight_ge(ws, col, first, last, threshold, fill="FBE3DE"):
    rng = f"{get_column_letter(col)}{first}:{get_column_letter(col)}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(
        operator=">=", formula=[str(threshold)],
        fill=PatternFill("solid", fgColor=fill)))


# ---------------------------------------------------------------- print setup

def print_ready(ws, ncols, repeat_rows="1:4", landscape=True):
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = repeat_rows


def sheet_defaults(ws, tab_color=None, zoom=100):
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.zoomScale = zoom
    ws.sheet_view.showGridLines = False
