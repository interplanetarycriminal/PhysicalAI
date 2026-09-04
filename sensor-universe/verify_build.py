"""Structural release gate for any built workbook.

    python3 verify_build.py esp32_sensor_universe_vNN.xlsx

Runs after patch_v50.py and before audit_workbook.py. Exit 1 on any failure.
See HANDOFF.md, operating rule 6. Every check here caught a real defect once.
"""
import csv
import itertools
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE / "data"))
sys.path.insert(0, str(HERE))

import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

import fusion  # noqa: E402
import loader  # noqa: E402
import outcome_solver as osv  # noqa: E402
import patch_v50  # noqa: E402
import solve  # noqa: E402
import vocab  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else str(patch_v50.OUT)
V50 = patch_v50.SRC
FIRST, LAST = 8, 412          # sensor rows on Kit Builder / Coverage Matrix
fails = []


def check(name, ok, detail=""):
    print(f"  {'✓' if ok else '✗ FAIL'} {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        fails.append(name)


# 1 · integrity + inventory
check("zip integrity", zipfile.ZipFile(OUT).testzip() is None)
wb = openpyxl.load_workbook(OUT)
if str(V50) != str(Path(OUT).resolve()) and Path(V50).exists():
    wb_src = openpyxl.load_workbook(V50)
    check("no source sheet lost", not (set(wb_src.sheetnames) - set(wb.sheetnames)))
owned = set(patch_v50.OWNED_SHEETS)
check("all owned sheets present", owned <= set(wb.sheetnames),
      str(owned - set(wb.sheetnames)) or f"{len(owned)} owned + {len(wb.sheetnames) - len(owned)} inherited")
check("no duplicated owned sheet (rebase artefact)",
      not any(n.rstrip("0123456789") in owned and n not in owned for n in wb.sheetnames))

# 2 · Sensor Catalog contiguity + filter
ws = wb["Sensor Catalog"]
pop = [r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value]
check("catalog contiguous from row 4", pop == list(range(4, 4 + len(pop))),
      f"{len(pop)} rows, last {pop[-1]}")
check("auto_filter spans the catalog",
      str(ws.auto_filter.ref) == f"$A$3:${get_column_letter(ws.max_column)}${pop[-1]}",
      str(ws.auto_filter.ref))

# 3 · Idea Forge wired to the whole catalog
forge = wb["Idea Forge"]
texts = []
for row in forge.iter_rows():
    for c in row:
        t = c.value.text if isinstance(c.value, ArrayFormula) else c.value
        if isinstance(t, str) and t.startswith("="):
            texts.append(t)
n_cat = len(pop)
check("Idea Forge: no stale ranges",
      not any("$241" in t or "RANDBETWEEN(1,238)" in t for t in texts))
check("Idea Forge: drivers span the catalog",
      sum(1 for t in texts if f"RANDBETWEEN(1,{n_cat})" in t) == 4)
check("Idea Forge: 41 formulas", len(texts) == 41, str(len(texts)))

# 4 · Coverage Matrix truth + alignment
records, _ = loader.load_all(persist_ids=False)
sol = osv.build(records, vocab.INFERENCE)
caps = fusion.matrix_capabilities()
cap_keys = {k for k, _kd in caps}
col_of = {k: 5 + i for i, (k, _kd) in enumerate(caps)}
cm, kb = wb["Coverage Matrix"], wb["Kit Builder"]
sensors = sorted((r for r in records if r.get("catalog") == "sensor"),
                 key=lambda r: int(r["id"][1:]))
mismatch = []
for i, r in enumerate(sensors):
    row = FIRST + i
    if cm.cell(row, 1).value != r["id"] or kb.cell(row, 1).value != r["id"]:
        mismatch.append((row, r["id"], "misaligned"))
        continue
    grid = {k for k in cap_keys if cm.cell(row, col_of[k]).value == 1}
    if grid != (fusion.covers(r) & cap_keys):
        mismatch.append((row, r["id"], "grid≠covers"))
check("matrix rows aligned + truthful", not mismatch and len(sensors) == LAST - FIRST + 1,
      str(mismatch[:2]) if mismatch else f"{len(sensors)} rows")

# 5 · formula-reference audit
rng = re.compile(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)")
inf_last = get_column_letter(4 + len(vocab.INFERENCE))
bad_refs, n_formulas = [], 0
for sheet in (cm, kb):
    for row in sheet.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_formulas += 1
                for m in rng.finditer(c.value):
                    r1, r2 = int(m.group(2)), int(m.group(4))
                    if (r1, r2) not in ((FIRST, LAST), (5, 5), (6, 6)) and r1 != r2 \
                            and r1 <= LAST:
                        bad_refs.append((sheet.title, c.coordinate, m.group(0)))
check("formula ranges sane", not bad_refs, str(bad_refs[:3]))
print(f"      ({n_formulas} live formulas)")
overrun = [c.coordinate for (c,) in kb.iter_rows(min_row=FIRST, max_row=LAST, min_col=5, max_col=5)
           if isinstance(c.value, str) and f"${inf_last}" not in c.value]
check("buys-next stops at inference columns", not overrun, str(overrun[:3]))

out_bad, found = [], 0
for row in kb.iter_rows(min_row=LAST + 1, max_row=kb.max_row, min_col=3, max_col=4):
    key_c, count_c = row
    if (key_c.value in vocab.INFERENCE and isinstance(count_c.value, str)
            and re.fullmatch(r"='Coverage Matrix'!\$[A-Z]+\$5", count_c.value)):
        found += 1
        if f"${get_column_letter(col_of[key_c.value])}$5" not in count_c.value:
            out_bad.append(key_c.value)
check("outcome block refs correct columns",
      found == len(vocab.INFERENCE) and not out_bad, f"{found}/{len(vocab.INFERENCE)} {out_bad[:2]}")

edge_bad = []
for (c,) in kb.iter_rows(min_row=LAST + 1, max_row=kb.max_row, min_col=4, max_col=4):
    v = c.value
    if isinstance(v, str) and v.startswith("=IF(AND") and '"UNLOCKED"' in v:
        edge_bad += [(c.coordinate, m.group(0)) for m in re.finditer(r"\$D\$(\d+)=", v)
                     if int(m.group(1)) >= c.row]
check("chained edges reference earlier rows only", not edge_bad, str(edge_bad[:3]))


# 6 · independent closure re-implementation
def closure2(recs):
    counts = Counter(c for r in recs for c in fusion.covers(r))
    fired, granted, changed = set(), set(), True
    while changed:
        changed = False
        for e in fusion.FUSION_EDGES:
            if e["key"] not in fired and all(
                    counts.get(c, 0) >= m or (m == 1 and c in granted) for c, m in e["requires"]):
                fired.add(e["key"]); granted.add(e["provides"]); changed = True
    return fired, granted


tiers_ok = True
for t, tc in zip(sol["tiers"], sol["fusion"]["tier_closures"]):
    ids = {s["id"] for s in sol["trace_cost"][:t["upto"]]}
    fired, granted = closure2([r for r in sensors if r["id"] in ids])
    em = len({g for g in granted if g in fusion.EMERGENT})
    if (len(fired), em) != (tc["fired"], tc["emergent"]):
        tiers_ok = False
check("closure independently reproduced per tier", tiers_ok)
preload = {s["id"] for s in sol["trace_cost"][:sol["tiers"][0]["upto"]]}
check("FOUNDATION preload written",
      all((kb.cell(FIRST + i, 4).value == 1) == (r["id"] in preload) for i, r in enumerate(sensors)),
      f"{len(preload)} preloaded")

# 7 · exports agree with the build
with open(HERE / "exports" / "coverage_matrix.csv") as f:
    rows = list(csv.reader(f))
hdr = rows[0][3:]
csv_ok = len(rows) == len(sensors) + 1 and all(
    rows[1 + i][0] == r["id"] and
    [int(x) for x in rows[1 + i][3:]] == [1 if k in fusion.covers(r) else 0 for k in hdr]
    for i, r in enumerate(sensors))
check("coverage_matrix.csv equals the grid", csv_ok)
with open(HERE / "exports" / "fusion_edges.csv") as f:
    check("fusion_edges.csv rows", sum(1 for _ in f) == len(fusion.FUSION_EDGES) + 1)
j = json.load(open(HERE / "exports" / "solver_run.json"))
check("solver_run.json fusion object", j.get("fusion", {}).get("n_edges") == len(fusion.FUSION_EDGES))


# 8 · best_groups sanity — the set ranker, re-derived rather than re-called
GR = sol.get("groups")
by_id = {r["id"]: r for r in records}
RANKINGS = [n for n, _o, _l in solve.GROUP_RANKINGS]
OBJ_OF = {n: o for n, o, _l in solve.GROUP_RANKINGS}
SHARED2 = {"I2C": (2, 0), "SPI": (3, 1), "1-Wire": (1, 0), "I2S": (3, 0)}


def pin_cost2(recs):
    """Independent re-implementation of the set pin model: a shared bus is paid
    once for the set then per device, and a multi-interface part takes whatever
    is cheapest given which buses the set has opened. Enumerated with
    combinations rather than a bitmask, so it is a genuine second opinion."""
    buses = sorted({i for r in recs for i in (r.get("iface") or []) if i in SHARED2})
    best = None
    for n in range(len(buses) + 1):
        for combo in itertools.combinations(buses, n):
            total, ok = sum(SHARED2[b][0] for b in combo), True
            for r in recs:
                ifc = r.get("iface") or []
                opts = [SHARED2[b][1] for b in combo if b in ifc]
                if "UART" in ifc:
                    opts.append(2)
                if not ifc or any(i not in SHARED2 and i != "UART" for i in ifc):
                    opts.append(int(r.get("pins") or 0))
                if not opts:
                    ok = False
                    break
                total += min(opts)
            if ok and (best is None or total < best):
                best = total
    return best if best is not None else sum(int(r.get("pins") or 0) for r in recs)


def fixed_addr2(r):
    """The one I2C address a part cannot be moved off, or None. Ranges are
    expanded, so '0x18-0x1F' is eight addresses and therefore not fixed."""
    if "I2C" not in (r.get("iface") or []) or len(r.get("iface") or []) > 1:
        return None
    t = r.get("i2c_addr") or ""
    got = set()
    for m in re.finditer(r"0x([0-9A-Fa-f]{2})\s*[-\u2013\u2014]\s*0x([0-9A-Fa-f]{2})", t):
        a, b = int(m.group(1), 16), int(m.group(2), 16)
        if a <= b and b - a <= 32:
            got |= set(range(a, b + 1))
    got |= {int(x, 16) for x in re.findall(r"0x([0-9A-Fa-f]{2})", t)}
    low = t.lower()
    if len(got) != 1 or any(w in low for w in osv.I2C_FLEX_WORDS):
        return None
    return got.pop()


check("groups block present", bool(GR) and all(k in GR for k in RANKINGS + ["groups"]))
GG = GR["groups"] if GR else []
kmax = GR["params"]["max_k"] if GR else 0

bad = [g["ids"] for g in GG
       if not (2 <= g["n_parts"] <= kmax)
       or len(set(g["ids"])) != g["n_parts"]
       or any(i not in by_id for i in g["ids"])]
check("groups are 2..k distinct catalog sensors", not bad,
      str(bad[:3]) if bad else f"{len(GG)} groups, k<={kmax}")

bad = [(g["ids"], g["usd"], round(sum(by_id[i].get("usd") or 0 for i in g["ids"]), 2))
       for g in GG if abs(g["usd"] - sum(by_id[i].get("usd") or 0 for i in g["ids"])) > 0.005]
check("group usd equals the sum of its members", not bad, str(bad[:3]) if bad else f"{len(GG)} ok")

bad = [(g["ids"], g["pins"], pin_cost2([by_id[i] for i in g["ids"]]))
       for g in GG if g["pins"] != pin_cost2([by_id[i] for i in g["ids"]])]
check("pin model independently reproduced", not bad,
      str(bad[:3]) if bad else f"{len(GG)} groups repriced")

bad = []
for g in GG:
    fixed = [a for a in (fixed_addr2(by_id[i]) for i in g["ids"]) if a is not None]
    if len(fixed) != len(set(fixed)):
        bad.append(g["ids"])
check("no group collides two rigid I2C addresses", not bad,
      str(bad[:3]) if bad else f"{len(GG)} groups clean")

# ...but that check alone has no teeth: every rigid collision scores far too low
# to reach a top-30 ranking, so GG is clean whether or not `outcome_solver`
# enforces the rule. Give it an oracle. Re-derive the colliding pairs from the
# catalog here, then make the ranker prove it rejects each one, on a two-part
# universe where no better set can crowd it out of the answer.
gp = GR["params"] if GR else {}
gpool = {}
for sid, (r, inf) in osv.index(records, vocab.INFERENCE).items():
    usd = r.get("usd") if isinstance(r.get("usd"), (int, float)) else 1e9
    if gp.get("max_usd") is not None and usd > gp["max_usd"]:
        continue
    if int(r.get("pins") or 0) > osv.MAX_PART_PINS or osv.esp32_unusable(r):
        continue
    gpool[sid] = (r, inf)

rigid = defaultdict(list)
for sid, (r, _inf) in gpool.items():
    a = fixed_addr2(r)
    if a is not None:
        rigid[a].append(sid)
rigid_pairs = sorted((a, x, y) for a, ids in rigid.items()
                     for x, y in itertools.combinations(sorted(ids), 2))
# a named pair to anchor the oracle: two thermal arrays welded to 0x33, one
# interface each, no jumper word anywhere in either address string
anchor = [(a, x, y) for a, x, y in rigid_pairs
          if "MLX90640" in gpool[x][0]["n"] and "MLX90641" in gpool[y][0]["n"]]
check("MLX90640 + MLX90641 re-derive as a rigid 0x33 collision",
      len(gpool) == gp.get("pool") and len(anchor) == 1 and anchor[0][0] == 0x33,
      f"pool {len(gpool)} vs {gp.get('pool')}, {len(rigid_pairs)} rigid pairs, "
      f"anchor {[(x, y) for _a, x, y in anchor]}")

survived = []
for a, x, y in rigid_pairs:
    res = osv.best_groups({x: gpool[x], y: gpool[y]},
                          records=[gpool[x][0], gpool[y][0]], min_k=2, max_k=2)
    if res["groups"] or res["params"]["rejected"] != 1:
        survived.append((f"0x{a:02x}", x, y, len(res["groups"]),
                         res["params"]["rejected"]))
check("the ranker really rejects every rigid I2C collision",
      bool(rigid_pairs) and not survived,
      str(survived[:3]) if survived else
      f"{len(rigid_pairs)} pairs, each rejected 1-for-1 on its own universe")

hit = [(name, g["ids"]) for name in RANKINGS for g in GR[name]
       for _a, x, y in anchor if x in g["ids"] and y in g["ids"]]
check("no ranking row contains the rigid 0x33 pair", not hit,
      str(hit[:3]) if hit else f"{sum(len(GR[n]) for n in RANKINGS)} rows scanned")

bad = []
for name in RANKINGS:
    vals = [g[OBJ_OF[name]] for g in GR[name]]
    if vals != sorted(vals, reverse=True):
        bad.append(name)
check("every ranking is monotone in its objective", not bad,
      str(bad) or f"{len(RANKINGS)} rankings")

bad = []
for name in RANKINGS:
    g = GR[name][0]
    recs_g = [by_id[i] for i in g["ids"]]
    fired, granted = closure2(recs_g)
    declared = len({k for r in recs_g for k in (r.get("inferences") or [])
                    if k in vocab.INFERENCE})
    em = len({x for x in granted if x in fusion.EMERGENT})
    if (declared, em, declared + em) != (g["declared"], g["emergent"], g["score"]):
        bad.append((name, g["ids"], declared, em, g["declared"], g["emergent"], g["score"]))
check("top row of each ranking recounted from scratch", not bad,
      str(bad[:2]) if bad else f"{len(RANKINGS)} top rows")

gcsv = HERE / "exports" / "best_groups.csv"
check("best_groups.csv exists", gcsv.exists())
if gcsv.exists():
    with open(gcsv) as f:
        grows = list(csv.reader(f))
    check("best_groups.csv header matches the writer", grows[0] == solve.GROUP_COLS,
          str(grows[0][:4]))
    want = sum(len(GR[n]) for n in RANKINGS)
    check("best_groups.csv row count equals the rankings", len(grows) - 1 == want,
          f"{len(grows) - 1} rows vs {want}")

print()
if fails:
    print(f"❌ {len(fails)} FAILURES: {fails}")
    sys.exit(1)
print(f"✅ all checks pass for {OUT}")
