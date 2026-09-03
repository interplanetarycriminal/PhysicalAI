#!/usr/bin/env python3
"""Audit ANY version of the Sensor Universe workbook for the errors that survive.

    python3 audit_workbook.py path/to/esp32_sensor_universe_v52.xlsx

Two genuine safety hazards and one hardware-destroying instruction survived
roughly forty-five iterations of this document because nothing was ever checking.
That is a process gap, not a content gap. This closes it: point it at any future
version and it re-runs every check that found something last time, plus the
general classes those errors belonged to.

Exit code is non-zero if anything CRITICAL is found, so it can gate a release.

Deliberately conservative: it flags patterns for a human to judge rather than
trying to be clever. A false positive costs ten seconds; a false negative cost
this document forty-five versions.
"""
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

CRITICAL, HIGH, MEDIUM, VERIFY = "CRITICAL", "HIGH", "MEDIUM", "VERIFY"

# ----------------------------------------------------------------- rules
# (severity, label, regex, why it matters, what to do)
RULES = [
 (CRITICAL, "Heated gas sensor sited inside a flammable volume",
  r"(MQ-?\d+|catalytic|pellistor|heated (?:bead|element))[^.]{0,200}"
  r"(gas locker|LPG locker|bottle storage|battery (?:room|shed|box)|fuel tank|bilge)"
  r"|(gas locker|LPG locker|battery (?:room|shed))[^.]{0,200}(MQ-?\d+|pellistor)",
  "An MQ or pellistor element is an unenclosed bead at ~300°C. In a volume where flammable gas "
  "can accumulate it is a candidate IGNITION SOURCE, not a safety device.",
  "Move it to the ventilated space OUTSIDE the volume, and point at a certified "
  "intrinsically-safe detector for the volume itself."),

 (CRITICAL, "Life-safety claim without a disclaimer",
  r"(prevents?|prevent(?:ing|ed)?|avoids?)[^.]{0,80}"
  r"(explosion|fire\b|fatal|poisoning|asphyxiation|carbon monoxide)",
  "Hobby sensors drift, poison and fail silently. A build framed as preventing a fatal event "
  "will be trusted as one.",
  "Add 'NOT a life-safety device — supplement, never replace, a certified alarm' to the entry."),

 (CRITICAL, "Output claimed safe for 3.3V logic while the supply exceeds it",
  r"(?<!not )(?<!isn.t )(?<!is NOT )(logic[- ]safe|3\.?3\s?V[- ]safe|safe (?:for|to)[^.]{0,30}(?:GPIO|logic))",
  "Open-collector and NPN industrial sensors swing to their SUPPLY rail. Documented as "
  "'logic-safe' at 12-24V, wiring one per the datasheet line destroys the GPIO.",
  "State the actual output swing, and whether a divider, level shifter or optocoupler is required."),

 (VERIFY, "Vendor part numbers cited — each is a checkable claim",
  r"\b(Adafruit|SparkFun|Pimoroni|Seeed)\s+\d{3,5}\b",
  "A specific SKU is a checkable claim. v5 cited 'Adafruit 5417' for an ADXL355 breakout that "
  "does not exist (5417 is a MicroPython Pyboard).",
  "Verify every numeric SKU against the vendor's site, or describe the board generically."),

 (VERIFY, "Library names cited — each is a checkable claim",
  r"\b(Adafruit|SparkFun)_[A-Za-z0-9]{3,}\b",
  "Library names follow a predictable pattern, which makes them easy to invent by accident. "
  "'Adafruit_ADXL355' was cited and does not exist.",
  "Check the vendor's GitHub org before citing a library."),

 (HIGH, "Absolute sensor described as gauge (or vice versa)",
  r"\bgauge pressure\b|\babsolute pressure\b",
  "An absolute sensor reads ~14.7 PSI on the bench. Every gauge use case (tank depth, "
  "tensiometer, cuff) silently breaks if the part is actually absolute.",
  "Confirm the part number's suffix. State explicitly whether atmospheric must be subtracted."),

 (HIGH, "Claim of immunity to a dominant error source",
  r"immune to (?:target )?colou?r|unaffected by (?:ambient|temperature|humidity)|"
  r"no calibration (?:required|needed)|never drifts?",
  "Blanket immunity claims are almost always false. ToF range depends strongly on target "
  "reflectivity; nearly everything drifts with temperature.",
  "Replace with the actual dependence and its magnitude."),

 (MEDIUM, "Marketing language where a specification belongs",
  r"MCERTS[- ]adjacent|research[- ]grade(?!\s*\()|lab[- ]grade(?!\s*\()|professional[- ]grade|"
  r"exceptional (?:sensitivity|accuracy|performance)|best[- ]in[- ]class",
  "Unquantified superlatives read as specifications and are not. 'MCERTS-adjacent' was used "
  "for a sensor that is not MCERTS certified.",
  "State the number, or state the certification, or say neither."),

 (MEDIUM, "Resistance/impedance range that may be inverted",
  r"(\d+(?:\.\d+)?)\s*[kKmM]?Ω\s*dark[^.]{0,40}(\d+(?:\.\d+)?)\s*[kKmM]?Ω\s*(?:bright|light)",
  "A GL5528 is ~1MΩ dark and ~10-20kΩ at 10 lux. v5 stated 10kΩ dark to 1kΩ bright — wrong "
  "by two orders of magnitude, which produces the wrong divider resistor.",
  "Check which way round the resistance goes, and give the value at a stated illuminance."),

 (MEDIUM, "Timing claim that ignores host latency",
  r"(nanosecond|microsecond|µs|ns)[^.]{0,60}(ESP32|GPIO|interrupt|Arduino)",
  "A PPS edge is ~50ns at the pin; an ESP32 interrupt timestamp is microseconds. Quoting the "
  "pin figure as achievable end-to-end is misleading.",
  "Quote the device figure and the achievable end-to-end figure separately."),

 (MEDIUM, "Sensor family sharing a fixed I2C address in a multi-sensor build",
  r"(VL53L\dC?X?|VL6180)[^.]{0,80}(×|x)\s?[2-9]|[2-9]\s?(×|x)\s?(VL53|VL6180)",
  "The whole ST VL53/VL6180 family is fixed at 0x29. Two on one bus is impossible without "
  "XSHUT sequencing or a multiplexer.",
  "Add the multiplexer or the XSHUT sequence to the build's parts list."),

 (MEDIUM, "Removed or deprecated SoC feature cited as available",
  r"hallRead|internal hall sensor|touchRead[^.]{0,60}(C3|C6|H2)|dacWrite[^.]{0,60}(C3|C6|H2|S3)",
  "The ESP32 internal hall sensor was removed in ESP-IDF v5. C3/C6/H2 have no touch peripheral "
  "and no DAC; S3 has no DAC.",
  "Check the variant's peripheral list before recommending it."),
]

# Sheets that legitimately discuss these patterns in order to warn about them.
EXEMPT_SHEETS = {"Corrections Log", "Failure Museum", "Illusions & Artifacts",
                 "The Anti-Catalog II",
                 "The Anti-Catalog", "Physics Cheatsheet", "I2C & Wiring Reality"}


def audit(path):
    wb = load_workbook(path, data_only=True)
    findings = defaultdict(list)
    cells = 0
    for ws in wb.worksheets:
        if ws.title in EXEMPT_SHEETS:
            continue
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or len(v) < 12:
                    continue
                cells += 1
                for sev, label, pat, why, fix in RULES:
                    m = re.search(pat, v, re.I)
                    if m:
                        findings[(sev, label, why, fix)].append(
                            (ws.title, cell.coordinate, v[max(0, m.start() - 60):m.end() + 90]))
    return wb, findings, cells


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = Path(sys.argv[1])
    wb, findings, cells = audit(path)

    print(f"\n  AUDIT · {path.name}")
    print(f"  {len(wb.sheetnames)} sheets · {cells:,} text cells scanned · "
          f"{len(RULES)} rules\n")

    order = {CRITICAL: 0, HIGH: 1, MEDIUM: 2, VERIFY: 3}
    n_crit = 0
    for (sev, label, why, fix), hits in sorted(findings.items(), key=lambda kv: order[kv[0][0]]):
        mark = {CRITICAL: "🔴", HIGH: "🟠", MEDIUM: "🟡", VERIFY: "🔎"}[sev]
        if sev == CRITICAL:
            n_crit += len(hits)
        print(f"{mark} {sev}  {label}   ({len(hits)} hit{'s' if len(hits) != 1 else ''})")
        print(f"     why: {why}")
        print(f"     fix: {fix}")
        for sheet, coord, excerpt in (hits[:2] if sev == VERIFY else hits[:4]):
            print(f"       · {sheet}!{coord}  …{' '.join(excerpt.split())[:150]}…")
        _shown = 2 if sev == VERIFY else 4
        if len(hits) > _shown:
            print(f"       · … and {len(hits) - _shown} more")
        print()

    if not findings:
        print("  ✓ nothing flagged\n")
    else:
        print(f"  {sum(len(v) for v in findings.values())} total flags, "
              f"{n_crit} critical\n")
        print("  These are patterns for a human to judge, not verdicts. A false positive costs\n"
              "  ten seconds; a false negative cost this document forty-five versions.\n")
    return 1 if n_crit else 0


if __name__ == "__main__":
    sys.exit(main())
